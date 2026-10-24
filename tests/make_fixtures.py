"""Generate deterministic test fixture tree in tests/_tree/ (Section 25.1)."""

from __future__ import annotations

import email.message
import email.policy
import shutil
import subprocess
from pathlib import Path


def create_fixture_tree(base_dir: Path | None = None) -> Path:
    """Create tests/_tree/ and populate it with sample documents, code, and config."""
    if base_dir is None:
        base_dir = Path(__file__).resolve().parent / "_tree"

    # Remove existing fixture tree if present
    if base_dir.exists():
        shutil.rmtree(base_dir, ignore_errors=True)

    base_dir.mkdir(parents=True, exist_ok=True)

    personal_dir = base_dir / "personal"
    personal_dir.mkdir(parents=True, exist_ok=True)
    (personal_dir / "car").mkdir(parents=True, exist_ok=True)

    legal_dir = base_dir / "legal"
    legal_dir.mkdir(parents=True, exist_ok=True)
    (legal_dir / "mail").mkdir(parents=True, exist_ok=True)

    code_dir = base_dir / "code"
    code_dir.mkdir(parents=True, exist_ok=True)
    repo_dir = code_dir / "demo_repo"
    repo_dir.mkdir(parents=True, exist_ok=True)
    (repo_dir / "app" / "auth").mkdir(parents=True, exist_ok=True)
    (repo_dir / "node_modules" / "x").mkdir(parents=True, exist_ok=True)
    (code_dir / "scripts").mkdir(parents=True, exist_ok=True)

    # -----------------------------------------------------------------------
    # 1. personal/
    # -----------------------------------------------------------------------
    import pymupdf

    # personal/Scan_0042.pdf (Passport)
    doc_passport = pymupdf.open()
    page1 = doc_passport.new_page()
    page1.insert_text(
        (50, 72),
        "REPUBLIC OF INDIA PASSPORT\n"
        "Passport No Z1234567\n"
        "Surname SHARMA Given Names RAVI\n"
        "Nationality INDIAN\n"
        "Date of Birth 15/08/1990\n"
        "Date of Issue 14/03/2019\n"
        "Date of Expiry 13/03/2029\n"
        "Place of Issue HYDERABAD",
        fontsize=12,
    )
    doc_passport.save(str(personal_dir / "Scan_0042.pdf"))
    doc_passport.close()

    # personal/car/policy_2026.pdf (Car insurance)
    doc_policy = pymupdf.open()
    page2 = doc_policy.new_page()
    page2.insert_text(
        (50, 72),
        "ICICI Lombard General Insurance Company\n"
        "Motor Insurance Certificate\n"
        "Policy No 3001/123456\n"
        "Period of Insurance 01/04/2026 to 31/03/2027\n"
        "Insured Name: Ravi Sharma\n"
        "Registration TS09 AB 1234\n"
        "Total Premium Rs. 14,200\n"
        "Sum Insured: Rs. 6,50,000",
        fontsize=12,
    )
    doc_policy.save(str(personal_dir / "car" / "policy_2026.pdf"))
    doc_policy.close()

    # personal/salary_slip_aug.pdf (Payslip)
    doc_slip = pymupdf.open()
    page3 = doc_slip.new_page()
    page3.insert_text(
        (50, 72),
        "Payslip for August 2026\n"
        "Employer Northwind Ltd\n"
        "Employee: Ravi Sharma\n"
        "Gross Salary 1,20,000\n"
        "Net Pay 98,500\n"
        "Payment Date: 31/08/2026",
        fontsize=12,
    )
    doc_slip.save(str(personal_dir / "salary_slip_aug.pdf"))
    doc_slip.close()

    # personal/notes.md
    (personal_dir / "notes.md").write_text(
        "# Personal Notes\n\n"
        "## Insurance\n"
        "Car policy renews before March 31, 2027.\n\n"
        "## Travel\n"
        "Passport expires 2029.\n",
        encoding="utf-8",
    )

    # personal/budget.csv
    (personal_dir / "budget.csv").write_text(
        "Category,Amount,Month\n"
        "Rent,45000,2026-08\n"
        "Groceries,15000,2026-08\n"
        "Utilities,6000,2026-08\n"
        "Total,66000,2026-08\n",
        encoding="utf-8",
    )

    # personal/big.bin (> 1 MB limit)
    (personal_dir / "big.bin").write_bytes(b"0" * (1200 * 1024))

    # personal/.env (Must be denied)
    (personal_dir / ".env").write_text("FAKE_TOKEN=abc123\n", encoding="utf-8")

    # personal/id_rsa (Must be denied)
    (personal_dir / "id_rsa").write_text("-----BEGIN RSA PRIVATE KEY-----\nFAKE\n-----END RSA PRIVATE KEY-----\n", encoding="utf-8")

    # -----------------------------------------------------------------------
    # 2. legal/
    # -----------------------------------------------------------------------
    import docx

    # legal/Acme MSA 2023.docx
    doc_msa = docx.Document()
    doc_msa.add_paragraph("Master Services Agreement")
    doc_msa.add_paragraph(
        "WHEREAS this Master Services Agreement is effective as of March 1, 2023 between "
        "Harborline Events LLC and Acme Facility Services Inc."
    )
    doc_msa.add_paragraph("Governing Law: Delaware.")
    doc_msa.add_paragraph("IN WITNESS WHEREOF, the parties hereto have executed this Agreement.")
    msa_path = legal_dir / "Acme MSA 2023.docx"
    doc_msa.save(str(msa_path))

    # legal/Acme Amendment 1.docx
    doc_amend = docx.Document()
    doc_amend.add_paragraph("Amendment No. 1 to the Master Services Agreement")
    doc_amend.add_paragraph(
        "Amendment No. 1 to the Master Services Agreement dated March 1, 2023 between "
        "Harborline Events LLC and ACME FACILITY SERVICES INC."
    )
    doc_amend.add_paragraph("WHEREAS effective June 1, 2026 the term is extended.")
    doc_amend.add_paragraph("IN WITNESS WHEREOF, the parties have signed.")
    doc_amend.save(str(legal_dir / "Acme Amendment 1.docx"))

    # legal/Acme MSA 2023 copy.docx (byte-identical duplicate)
    shutil.copyfile(msa_path, legal_dir / "Acme MSA 2023 copy.docx")

    # legal/mail/signed.eml (Email with attached PDF)
    # Generate attachment PDF
    doc_att = pymupdf.open()
    att_page = doc_att.new_page()
    att_page.insert_text(
        (50, 72),
        "Executed Copy of Amendment No. 1\n"
        "Between Harborline Events LLC and Acme Facility Services Inc\n"
        "Effective June 1, 2026\n",
        fontsize=12,
    )
    att_bytes = doc_att.tobytes()
    doc_att.close()

    msg = email.message.EmailMessage()
    msg["Subject"] = "Signed Amendment No. 1"
    msg["From"] = "contracts@acme.com"
    msg["To"] = "ravi@harborline.com"
    msg["Date"] = "Mon, 01 Jun 2026 10:00:00 +0000"
    msg.set_content("Please find attached the signed Amendment No. 1.\nBest regards,\nAcme Legal Team")
    msg.add_attachment(
        att_bytes,
        maintype="application",
        subtype="pdf",
        filename="contract.pdf",
    )

    with open(legal_dir / "mail" / "signed.eml", "wb") as f:
        f.write(msg.as_bytes())

    # -----------------------------------------------------------------------
    # 3. code/
    # -----------------------------------------------------------------------
    # demo_repo: init git repo if git is available
    git_bin = shutil.which("git")
    if git_bin:
        try:
            subprocess.run([git_bin, "init"], cwd=str(repo_dir), capture_output=True, check=True)
            subprocess.run([git_bin, "config", "user.name", "Test User"], cwd=str(repo_dir), capture_output=True)
            subprocess.run([git_bin, "config", "user.email", "test@example.com"], cwd=str(repo_dir), capture_output=True)
        except Exception:
            (repo_dir / ".git").mkdir(parents=True, exist_ok=True)
    else:
        (repo_dir / ".git").mkdir(parents=True, exist_ok=True)

    (repo_dir / "README.md").write_text("# Demo App\nSample repo for FileGraph testing.\n", encoding="utf-8")
    (repo_dir / "pyproject.toml").write_text('[project]\nname = "demo-app"\nversion = "0.1.0"\n', encoding="utf-8")
    (repo_dir / "app" / "__init__.py").write_text('"""App package."""\n', encoding="utf-8")

    (repo_dir / "app" / "auth" / "session.py").write_text(
        '"""User authentication and session handling."""\n\n'
        'class Session:\n'
        '    """User session container."""\n'
        '    def __init__(self, token: str):\n'
        '        self.token = token\n\n'
        'def login_user(username: str, password: str) -> Session:\n'
        '    """Log in user and return session."""\n'
        '    return Session(token="tok-123")\n\n'
        'def verify_token(token: str) -> bool:\n'
        '    """Verify validity of token."""\n'
        '    return bool(token)\n',
        encoding="utf-8",
    )

    (repo_dir / "app" / "api.py").write_text(
        'from app.auth.session import login_user\n\n'
        'def handle_login(req: dict) -> dict:\n'
        '    sess = login_user(req["user"], req["pass"])\n'
        '    return {"token": sess.token}\n',
        encoding="utf-8",
    )

    # config.py has a fake secret that must be redacted
    (repo_dir / "config.py").write_text(
        'API_KEY = "sk-test-0123456789abcdefghij"\n'
        'DEBUG = True\n',
        encoding="utf-8",
    )

    # node_modules/x/index.js (Must be ignored)
    (repo_dir / "node_modules" / "x" / "index.js").write_text(
        'console.log("node_modules test");\n', encoding="utf-8"
    )

    # scripts/tool.js
    (code_dir / "scripts" / "tool.js").write_text(
        'export function signIn() {}\n', encoding="utf-8"
    )

    # -----------------------------------------------------------------------
    # 4. tests/_tree/filegraph.yaml
    # -----------------------------------------------------------------------
    yaml_content = f"""me:
  names: ["Ravi Sharma"]
  organizations: ["Harborline Events LLC"]

roots:
  - path: {personal_dir.as_posix()}
    name: personal
    profile: auto
    llm: cloud
  - path: {legal_dir.as_posix()}
    name: legal
    profile: auto
    llm: cloud
  - path: {code_dir.as_posix()}
    name: code
    profile: auto
    llm: cloud

ignore:
  - "node_modules/"
  - ".git/"
  - ".venv/"
  - "__pycache__/"

deny:
  - "**/.ssh/**"
  - "**/*.kdbx"

limits:
  max_file_mb: 1
  max_text_chars_per_file: 2000000
"""
    (base_dir / "filegraph.yaml").write_text(yaml_content, encoding="utf-8")

    return base_dir


if __name__ == "__main__":
    p = create_fixture_tree()
    print(f"Fixture tree created at {p}")
